#!/usr/bin/env python3
"""
Extract the DMI use-case catalogue from the consolidated TNO source document
(source-docs/<SOURCE>) into src/data/usecases.ts.

Run with `npm run extract` after any revision of the Word document.

The document is one report with a chapter per domain ("Use cases voor het
energiedomein", ...). Every use case is a two-column table with a consistent
9-row schema (naam / hoofd thema / waarom / niveau / beleidsdomeinen /
functionaliteiten / analyse / data / voorbeelden). The document body is
walked in order so that every table inherits the H1/H2/H3 heading context
printed above it.

The chapter on ruimtelijke ordening & wonen contains no use-case tables; it
describes a five-step approach instead and is emitted separately.
"""

from __future__ import annotations

import json
import re
import sys
import unicodedata
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

try:
    import docx
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph
except ImportError:  # pragma: no cover
    sys.exit("python-docx is required:  pip3 install python-docx")

ROOT = Path(__file__).resolve().parent.parent
DOCS = ROOT / "source-docs"
OUT = ROOT / "src" / "data" / "usecases.ts"

SOURCE = "TNO-2026-17923 Use cases catalogus.docx"
# The report has no "vastgesteld door" on its title page and its introduction
# is marked as a concept version.
STATUS = "concept"

# The italic placeholder text that marks the blank format table reproduced in
# the "Uitgangspunt" chapter.
TEMPLATE_MARKER = "aansprekende en duidelijke naam"


# --------------------------------------------------------------------------
# Chapters → domains
# --------------------------------------------------------------------------


@dataclass(frozen=True)
class Chapter:
    # Lower-cased fragment of the H1 that identifies the chapter.
    keyword: str
    domain: str
    # Most chapters nest use cases (H3) under a thema (H2); gezondheid lists
    # its use cases straight under the chapter, with one mis-styled H2.
    has_thema_level: bool = True
    # The ruimtelijke ordening chapter has no use-case tables at all.
    is_aanpak: bool = False


CHAPTERS = [
    Chapter("energie", "energie"),
    Chapter("verkeer", "mobiliteit"),
    Chapter("mobiliteit", "mobiliteit"),
    Chapter("circulariteit", "circulariteit"),
    Chapter("klimaatadaptatie", "klimaatadaptatie"),
    Chapter("gezondheid", "gezondheid", has_thema_level=False),
    Chapter("veiligheid", "veiligheid"),
    Chapter("milieu", "milieu_natuur_ecologie"),
    Chapter("ruimtelijke ordening", "ruimtelijke_ordening_wonen", is_aanpak=True),
]

DOMAIN_ORDER = [c.domain for c in CHAPTERS if not c.is_aanpak]
DOMAIN_ORDER = list(dict.fromkeys(DOMAIN_ORDER))


def chapter_for(h1: str) -> Chapter | None:
    low = h1.lower()
    for chapter in CHAPTERS:
        if chapter.keyword in low:
            return chapter
    return None


# --------------------------------------------------------------------------
# Text helpers
# --------------------------------------------------------------------------


# Word stores Symbol/Wingdings list bullets as Private Use Area code points
# ( is the Symbol-font "•"). They have no glyph in a web font and would
# render as tofu, so they are dropped rather than carried through.
PRIVATE_USE = re.compile("[\\ue000-\\uf8ff]")


def norm(text: str) -> str:
    """NFKC-normalise, drop Word's symbol-font bullets, tidy hyphens/nbsp."""
    text = PRIVATE_USE.sub("", unicodedata.normalize("NFKC", text))
    return text.replace("\xa0", " ").replace("‐", "-").replace("‑", "-").strip()


def clean_heading(text: str) -> str:
    """Strip Word numbering ("2.4.2 "), stray bold markers and doubled words."""
    text = norm(text).replace("**", "")
    text = re.sub(r"^\d+(\.\d+)*\.?\s+", "", text)
    text = re.sub(r"\s+", " ", text).strip(" .")
    # "Warmte-koude conversie conversie en opslag" — an editing leftover.
    return re.sub(r"\b(\w+)(\s+\1\b)+", r"\1", text)


def clean_thema(text: str) -> str:
    """"Nieuwe mobiliteitsvormen Thema" → "Nieuwe mobiliteitsvormen"."""
    return re.sub(r"\s+thema$", "", text, flags=re.I).strip()


BULLETS = "•·‣▪◦-–—*"


def to_lines(cell_text: str) -> list[str]:
    """Split a cell into list items, stripping bullet glyphs and empties."""
    out: list[str] = []
    for raw in norm(cell_text).split("\n"):
        line = raw.strip().lstrip(BULLETS).strip()
        if line:
            out.append(line)
    return out


def to_prose(cell_text: str) -> str:
    return "\n".join(to_lines(cell_text))


def first_sentence(prose: str) -> str:
    """A card-sized teaser: grow by sentence until it reads as a full thought."""
    flat = " ".join(prose.split())
    if not flat:
        return ""
    parts = re.split(r"(?<=[.!?])\s+", flat)
    teaser = ""
    for part in parts:
        teaser = f"{teaser} {part}".strip()
        if len(teaser) >= 90:
            break
    return teaser


def slugify(name: str) -> str:
    ascii_name = (
        unicodedata.normalize("NFKD", name).encode("ascii", "ignore").decode("ascii")
    )
    slug = re.sub(r"[^a-z0-9]+", "-", ascii_name.lower()).strip("-")
    words = slug.split("-")
    while len("-".join(words)) > 60 and len(words) > 3:
        words.pop()
    return "-".join(words)


# --------------------------------------------------------------------------
# Field label matching (labels vary in spelling and casing between chapters)
# --------------------------------------------------------------------------

FIELD_PREFIXES = {
    "naam": ("use case naam", "use case name"),
    "hoofdthema": ("hoofd thema", "hoofdthema"),
    "waarom": ("beschrijving",),
    "niveau": ("niveau",),
    "beleidsdomeinen": ("relevante beleidsdomeinen",),
    "functionaliteiten": ("benodigde functionaliteiten",),
    "analyse": ("benodigde analyse",),
    "data": ("benodigde data",),
    "voorbeelden": ("voorbeeld",),
}


def field_key(label: str) -> str | None:
    label = norm(label).replace("\n", " ").lower()
    for key, prefixes in FIELD_PREFIXES.items():
        if any(label.startswith(p) for p in prefixes):
            return key
    return None


def read_table(table: Table) -> dict[str, str]:
    """Map a use-case table onto our canonical field keys."""
    fields: dict[str, str] = {}
    rows = table.rows
    for index, row in enumerate(rows):
        cells = row.cells
        if len(cells) < 2:
            continue
        key = field_key(cells[0].text)
        if key is None:
            continue
        value = cells[1].text
        if key not in fields:
            fields[key] = value
        elif index == len(rows) - 1 and "voorbeelden" not in fields:
            # One table repeats the "functionaliteiten" label on its final
            # row; by position and content that row holds the voorbeelden.
            fields["voorbeelden"] = value
    return fields


# --------------------------------------------------------------------------
# Taxonomy normalisation
# --------------------------------------------------------------------------

# Keyword → one of the canonical 8 beleidsdomeinen. A single raw token may hit
# several (e.g. "Wonen en Ruimtelijke Ordening"), which is intentional.
BELEIDSDOMEIN_KEYWORDS: list[tuple[str, str]] = [
    ("ruimtelijke ordening", "ruimtelijke_ordening"),
    ("ruimtelijke ontwikkeling", "ruimtelijke_ordening"),
    ("ruimtelijke planning", "ruimtelijke_ordening"),
    ("gebouwde omgeving", "ruimtelijke_ordening"),
    ("spatial planning", "ruimtelijke_ordening"),
    (r"\bro\b", "ruimtelijke_ordening"),
    ("wonen", "wonen"),
    ("housing", "wonen"),
    ("verkeer", "verkeer_vervoer"),
    ("mobiliteit", "verkeer_vervoer"),
    ("mobility", "verkeer_vervoer"),
    ("transport", "verkeer_vervoer"),
    ("logistiek", "verkeer_vervoer"),
    ("scheepvaart", "verkeer_vervoer"),
    ("bereikbaarheid", "verkeer_vervoer"),
    ("veiligheid", "veiligheid"),
    ("safety", "veiligheid"),
    ("energie", "energie_klimaat"),
    ("energy", "energie_klimaat"),
    ("klimaat", "energie_klimaat"),
    ("climate", "energie_klimaat"),
    ("laadinfra", "energie_klimaat"),
    ("laden", "energie_klimaat"),
    ("netbeheer", "energie_klimaat"),
    ("warmte", "energie_klimaat"),
    ("duurzaamheid", "energie_klimaat"),
    ("luchtkwaliteit", "energie_klimaat"),
    ("milieu", "milieu_natuur_ecologie"),
    ("environment", "milieu_natuur_ecologie"),
    ("natuur", "milieu_natuur_ecologie"),
    ("nature", "milieu_natuur_ecologie"),
    ("ecolog", "milieu_natuur_ecologie"),
    ("biodivers", "milieu_natuur_ecologie"),
    ("groen", "milieu_natuur_ecologie"),
    ("water", "milieu_natuur_ecologie"),
    ("gezondheid", "gezondheid_welzijn"),
    ("health", "gezondheid_welzijn"),
    ("welzijn", "gezondheid_welzijn"),
    ("zorg", "gezondheid_welzijn"),
    ("sociaal domein", "gezondheid_welzijn"),
    ("inclusie", "gezondheid_welzijn"),
    ("economie", "economische_vitaliteit"),
    ("economisch", "economische_vitaliteit"),
    ("economic", "economische_vitaliteit"),
    ("circulari", "economische_vitaliteit"),
    ("circular", "economische_vitaliteit"),
    ("bedrijventerrein", "economische_vitaliteit"),
    ("innovatie", "economische_vitaliteit"),
]

BELEIDSDOMEIN_ORDER = [
    "ruimtelijke_ordening",
    "wonen",
    "verkeer_vervoer",
    "veiligheid",
    "energie_klimaat",
    "milieu_natuur_ecologie",
    "gezondheid_welzijn",
    "economische_vitaliteit",
]


def split_domains(raw: str) -> list[str]:
    flat = norm(raw).replace("\n", ";")
    return [p.strip(" .()") for p in re.split(r"[;,/]", flat) if p.strip(" .()")]


def parse_beleidsdomeinen(raw: str) -> tuple[list[str], list[str]]:
    """Return (canonical 8, everything else kept verbatim)."""
    canonical: list[str] = []
    overige: list[str] = []
    for token in split_domains(raw):
        low = token.lower()
        hits = [
            key
            for kw, key in BELEIDSDOMEIN_KEYWORDS
            if (re.search(kw, low) if kw.startswith(r"\b") else kw in low)
        ]
        if hits:
            for key in hits:
                if key not in canonical:
                    canonical.append(key)
        else:
            cleaned = re.sub(r"^alle\s*", "", token, flags=re.I).strip()
            cleaned = cleaned[:1].upper() + cleaned[1:]
            if cleaned and cleaned not in overige:
                overige.append(cleaned)
    canonical.sort(key=BELEIDSDOMEIN_ORDER.index)
    return canonical, overige


NIVEAU_ORDER = ["strategisch", "tactisch", "operationeel"]
NIVEAU_ALIASES = {
    "strategisch": ("strateg",),
    "tactisch": ("tactis", "tactic"),
    "operationeel": ("operation",),
}


def parse_niveaus(raw: str) -> list[str]:
    low = norm(raw).lower()
    if "alle niveau" in low or "all levels" in low:
        return list(NIVEAU_ORDER)
    return [n for n in NIVEAU_ORDER if any(a in low for a in NIVEAU_ALIASES[n])]


# --------------------------------------------------------------------------
# Document walking
# --------------------------------------------------------------------------


@dataclass
class Line:
    text: str
    is_list: bool


@dataclass
class Extracted:
    fields: dict[str, str]
    domain: str
    thema: str
    heading: str
    toelichting: str = ""
    # Prose printed *after* the table (H4/H5 sections), kept as titled blocks.
    verdieping: list[dict] = field(default_factory=list)


@dataclass
class ThemaRecord:
    id: str
    domain: str
    name: str
    intro: str = ""
    usecase_ids: list[str] = field(default_factory=list)


@dataclass
class Walked:
    tables: list[Extracted]
    # (domain, thema) → intro lines
    thema_intros: dict[tuple[str, str], list[Line]]
    domain_intros: dict[str, list[Line]]
    # Prose before the first chapter: the "Uitgangspunt".
    uitgangspunt: list[str]
    # The blank format table explaining each field.
    formaat: list[tuple[str, str]]
    # The ruimtelijke ordening & wonen chapter: intro + five steps.
    aanpak: dict


def iter_blocks(document):
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield Table(child, document)


CHAPTER_TITLE = re.compile(r"^use cases? voor", re.I)


def heading_level(paragraph: Paragraph) -> int | None:
    style = paragraph.style.name or ""
    if style.startswith("Heading"):
        tail = style.split()[-1]
        return int(tail) if tail.isdigit() else 1
    # The mobility chapter title carries a custom style instead of Heading 1.
    if CHAPTER_TITLE.match(norm(paragraph.text)):
        return 1
    return None


FRONT_NOISE = ("alle rechten voorbehouden", "niets uit deze uitgave", "figuur", "inhoudsopgave")


def join_lines(lines: list[Line]) -> str:
    return "\n".join(line.text for line in lines)


def strip_subcase_list(lines: list[Line]) -> list[Line]:
    """Drop "Dit thema bestaat uit vier sub-use cases:" and its bullet list —
    the catalogue lists the use cases itself."""
    out: list[Line] = []
    skipping = False
    for line in lines:
        if re.search(r"bestaat uit .*sub-?\s?use ?cases?:?$", line.text, re.I):
            skipping = True
            continue
        if skipping and line.is_list:
            continue
        skipping = False
        out.append(line)
    return out


def walk(path: Path) -> Walked:
    document = docx.Document(str(path))
    headings: dict[int, str] = {}
    chapter: Chapter | None = None
    in_front = True

    tables: list[Extracted] = []
    thema_intros: dict[tuple[str, str], list[Line]] = {}
    domain_intros: dict[str, list[Line]] = {}
    uitgangspunt: list[str] = []
    formaat: list[tuple[str, str]] = []
    aanpak: dict = {"intro": [], "stappen": []}
    thema_closed: set[tuple[str, str]] = set()

    # Prose under the current use-case heading, before its table.
    pending: list[Line] = []
    # Whether the current use-case heading has produced a table yet.
    heading_has_table = False

    def usecase_heading() -> str:
        assert chapter is not None
        if chapter.has_thema_level:
            return headings.get(3, "")
        return headings.get(3) or headings.get(2, "")

    def current_thema() -> str:
        assert chapter is not None
        return clean_thema(headings.get(2, "")) if chapter.has_thema_level else ""

    def flush_pending() -> None:
        """A use-case heading that never got a table: its prose introduces
        the thema (energie's "Flexibiliteit") or is a step (ruimtelijke
        ordening)."""
        nonlocal pending
        if chapter is None or not pending:
            pending = []
            return
        if chapter.is_aanpak:
            aanpak["stappen"].append(
                {"naam": headings.get(3, ""), "toelichting": join_lines(pending)}
            )
        elif not heading_has_table:
            thema = current_thema()
            if thema:
                thema_intros.setdefault((chapter.domain, thema), []).extend(pending)
        pending = []

    for block in iter_blocks(document):
        if isinstance(block, Paragraph):
            style = block.style.name or ""
            level = heading_level(block)
            if level is not None:
                text = clean_heading(block.text)
                if level <= 3:
                    flush_pending()
                    heading_has_table = False
                elif chapter is not None and not chapter.is_aanpak and tables and heading_has_table:
                    # H4/H5 after a table opens a titled block of verdieping.
                    # Long "headings" are mis-styled bullets and stay prose.
                    blocks = tables[-1].verdieping
                    if len(text) > 80:
                        if not blocks:
                            blocks.append({"kop": "", "tekst": []})
                        blocks[-1]["tekst"].append(text)
                    else:
                        blocks.append({"kop": text, "tekst": []})
                    continue
                headings[level] = text
                for deeper in [k for k in headings if k > level]:
                    del headings[deeper]
                if level == 1:
                    in_front = False
                    chapter = chapter_for(text)
                    if chapter is not None:
                        domain_intros.setdefault(chapter.domain, [])
                elif level == 2 and chapter is not None and chapter.has_thema_level:
                    thema_intros.setdefault((chapter.domain, clean_thema(text)), [])
                continue

            body = norm(block.text)
            if not body or style.lower().startswith("toc"):
                continue
            line = Line(body, style.lower().startswith("opsomming") or style.lower().startswith("list"))

            if in_front:
                if not body.lower().startswith(FRONT_NOISE):
                    uitgangspunt.append(body)
                continue
            if chapter is None or body.lower().startswith("figuur"):
                continue

            h2, h3 = headings.get(2), headings.get(3)
            if chapter.is_aanpak:
                (pending if h3 else aanpak["intro"]).append(line)
                continue
            if h2 is None:
                domain_intros[chapter.domain].append(line)
            elif chapter.has_thema_level and h3 is None:
                key = (chapter.domain, clean_thema(h2))
                if key not in thema_closed:
                    thema_intros.setdefault(key, []).append(line)
            elif heading_has_table and tables:
                blocks = tables[-1].verdieping
                if not blocks:
                    blocks.append({"kop": "", "tekst": []})
                blocks[-1]["tekst"].append(body)
            else:
                pending.append(line)
            continue

        # ---- table ----
        fields = read_table(block)
        if "naam" not in fields:
            if fields and tables and heading_has_table:
                # A row split off into its own table (gezondheidsgelijkheid's
                # "benodigde data") continues the previous use case.
                for key, value in fields.items():
                    tables[-1].fields.setdefault(key, value)
            pending = []
            continue
        if TEMPLATE_MARKER in norm(fields["naam"]).lower():
            if not formaat:
                formaat = [
                    (norm(row.cells[0].text).replace("\n", " / "), to_prose(row.cells[1].text))
                    for row in block.rows
                    if len(row.cells) > 1 and norm(row.cells[0].text)
                ]
            pending = []
            continue
        if chapter is None or chapter.is_aanpak:
            pending = []
            continue
        thema = current_thema()
        tables.append(
            Extracted(
                fields=fields,
                domain=chapter.domain,
                thema=thema,
                heading=usecase_heading(),
                toelichting=join_lines(pending),
            )
        )
        if thema:
            thema_closed.add((chapter.domain, thema))
        pending = []
        heading_has_table = True

    flush_pending()

    return Walked(
        tables=tables,
        thema_intros=thema_intros,
        domain_intros=domain_intros,
        uitgangspunt=uitgangspunt,
        formaat=formaat,
        aanpak=aanpak,
    )


# --------------------------------------------------------------------------
# Report metadata from the title page (text boxes python-docx does not walk)
# --------------------------------------------------------------------------


def title_page(path: Path) -> dict[str, str]:
    with zipfile.ZipFile(path) as archive:
        xml = archive.read("word/document.xml").decode("utf-8")
    text = " ".join(re.findall(r"<w:t[^>]*>([^<]*)</w:t>", xml))
    text = re.sub(r"\s+", " ", text)
    kenmerk = re.search(r"\bTNO-\d{4}-\d+\b", text)
    datum = re.search(r"\b(\d{1,2} \w+ \d{4})\b", text)
    document = docx.Document(str(path))
    return {
        "titel": document.core_properties.title or "",
        "kenmerk": kenmerk.group(0) if kenmerk else "",
        "datum": datum.group(1) if datum else "",
    }


# --------------------------------------------------------------------------
# Build records
# --------------------------------------------------------------------------


def choose_name(heading: str, naam_cell: str) -> str:
    """The heading is authoritative; a few headings are full sentences while
    the naam cell holds the actual title, and a few tables have no heading."""
    name = heading or naam_cell
    if len(heading) > 90 and 0 < len(naam_cell) <= 90:
        name = naam_cell
    return name.strip(" .")


def build(path: Path):
    walked = walk(path)
    usecases: list[dict] = []
    themas: dict[tuple[str, str], ThemaRecord] = {}
    seen_slugs: set[str] = set()
    missing_fields: list[tuple[str, list[str]]] = []

    for item in walked.tables:
        fields = item.fields
        naam_cell = to_prose(fields.get("naam", ""))
        name = choose_name(item.heading, naam_cell)

        absent = [k for k in FIELD_PREFIXES if k not in fields]
        if absent:
            missing_fields.append((name, absent))

        waarom = to_prose(fields.get("waarom", ""))
        niveau_raw = to_prose(fields.get("niveau", ""))
        domeinen_raw = to_prose(fields.get("beleidsdomeinen", ""))
        canonical, overige = parse_beleidsdomeinen(domeinen_raw)

        slug = slugify(name)
        suffix = 2
        while slug in seen_slugs:
            slug = f"{slugify(name)}-{suffix}"
            suffix += 1
        seen_slugs.add(slug)

        if item.thema:
            key = (item.domain, item.thema)
            if key not in themas:
                intro = strip_subcase_list(walked.thema_intros.get(key, []))
                themas[key] = ThemaRecord(
                    id=slugify(f"{item.domain}-{item.thema}"),
                    domain=item.domain,
                    name=item.thema,
                    intro=join_lines(intro),
                )
            themas[key].usecase_ids.append(slug)

        usecases.append(
            {
                "id": slug,
                "name": name,
                "domain": item.domain,
                "thema": item.thema,
                "hoofdthema": to_prose(fields.get("hoofdthema", "")),
                "samenvatting": first_sentence(waarom),
                "toelichting": item.toelichting,
                "waarom": waarom,
                "niveaus": parse_niveaus(niveau_raw),
                "niveauRaw": niveau_raw,
                "beleidsdomeinen": canonical,
                "overigeBeleidsdomeinen": overige,
                "beleidsdomeinenRaw": domeinen_raw,
                "functionaliteiten": to_lines(fields.get("functionaliteiten", "")),
                "analyse": to_lines(fields.get("analyse", "")),
                "benodigdeData": to_lines(fields.get("data", "")),
                "voorbeelden": to_lines(fields.get("voorbeelden", "")),
                "verdieping": [
                    block for block in item.verdieping if block["kop"] or block["tekst"]
                ],
                "bron": {"document": SOURCE, "status": STATUS},
            }
        )

    # Keep the catalogue in chapter order even where a domain's chapters are
    # split (none are today, but the walk does not assume it).
    usecases.sort(key=lambda uc: DOMAIN_ORDER.index(uc["domain"]))

    domain_intros = {
        domain: join_lines(walked.domain_intros.get(domain, [])) for domain in DOMAIN_ORDER
    }
    bronnen = [
        {
            "domain": domain,
            "document": SOURCE,
            "status": STATUS,
            "aantal": sum(1 for uc in usecases if uc["domain"] == domain),
        }
        for domain in DOMAIN_ORDER
    ]
    aanpak = {
        "intro": join_lines(walked.aanpak["intro"]),
        "stappen": walked.aanpak["stappen"],
    }
    return (
        usecases,
        list(themas.values()),
        domain_intros,
        "\n".join(walked.uitgangspunt),
        walked.formaat,
        bronnen,
        aanpak,
        missing_fields,
    )


# --------------------------------------------------------------------------
# Emit TypeScript
# --------------------------------------------------------------------------


def ts(value) -> str:
    return json.dumps(value, ensure_ascii=False, indent=2)


HEADER = """// ============================================================
// AUTO-GENERATED — do not edit.
//
// Generated by scripts/extract-usecases.py from the consolidated DMI
// use-case catalogue (TNO report). Re-run `npm run extract` after any
// revision of that document.
//
// Labels, colours and display taxonomy live in ./taxonomy.ts and are
// maintained by hand.
// ============================================================

import type {
  Aanpak,
  Bron,
  BronDocument,
  Domain,
  FormaatVeld,
  Thema,
  UseCase,
} from "./types";

"""


def main() -> None:
    path = DOCS / SOURCE
    if not path.exists():
        sys.exit(f"Missing source document: {path.relative_to(ROOT)}")

    (
        usecases,
        themas,
        domain_intros,
        uitgangspunt,
        formaat,
        bronnen,
        aanpak,
        missing_fields,
    ) = build(path)
    meta = title_page(path)

    thema_json = [
        {
            "id": t.id,
            "domain": t.domain,
            "name": t.name,
            "intro": t.intro,
            "usecaseIds": t.usecase_ids,
        }
        for t in themas
    ]

    body = HEADER
    body += f"export const useCases: UseCase[] = {ts(usecases)};\n\n"
    body += f"export const themas: Thema[] = {ts(thema_json)};\n\n"
    body += (
        "export const domainIntros: Record<Domain, string> = "
        f"{ts(domain_intros)};\n\n"
    )
    body += f"export const uitgangspunt = {ts(uitgangspunt)};\n\n"
    body += (
        "export const formaat: FormaatVeld[] = "
        f"{ts([{'veld': v, 'toelichting': t} for v, t in formaat])};\n\n"
    )
    body += (
        "export const bronnen: (Bron & { domain: Domain; aantal: number })[] = "
        f"{ts(bronnen)};\n\n"
    )
    body += (
        "export const bronDocument: BronDocument = "
        f"{ts({'document': SOURCE, 'status': STATUS, **meta})};\n\n"
    )
    body += f"export const ruimtelijkeOrdeningWonen: Aanpak = {ts(aanpak)};\n"

    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(body, encoding="utf-8")

    # Report — the counts are the contract with the verification step.
    by_domain: dict[str, int] = {}
    for uc in usecases:
        by_domain[uc["domain"]] = by_domain.get(uc["domain"], 0) + 1
    print(f"Wrote {OUT.relative_to(ROOT)}")
    print(f"  bron: {meta['titel']} ({meta['kenmerk']}, {meta['datum']})")
    print(f"  {len(usecases)} use cases across {len(by_domain)} domeinen")
    for domain, count in by_domain.items():
        print(f"    {domain:<24} {count}")
    print(f"  {len(themas)} thema's")
    print(f"  {len(aanpak['stappen'])} stappen ruimtelijke ordening & wonen")
    overig = sorted({o for uc in usecases for o in uc["overigeBeleidsdomeinen"]})
    if overig:
        print("  beleidsdomeinen buiten de canonieke 8 (behouden, niet weggegooid):")
        for o in overig:
            print(f"    - {o}")
    empty = [uc["id"] for uc in usecases if not uc["beleidsdomeinen"]]
    if empty:
        print(f"  WARNING: no beleidsdomein matched for: {', '.join(empty)}")
    no_niveau = [uc["id"] for uc in usecases if not uc["niveaus"]]
    if no_niveau:
        print(f"  WARNING: no niveau matched for: {', '.join(no_niveau)}")
    if missing_fields:
        print("  WARNING: tables missing fields:")
        for name, absent in missing_fields:
            print(f"    - {name[:60]}: {', '.join(absent)}")


if __name__ == "__main__":
    main()
